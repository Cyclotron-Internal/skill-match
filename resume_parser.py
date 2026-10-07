"""Read plain text out of a resume file (.pdf, .docx, .txt)."""
from pathlib import Path

from docx import Document
from pypdf import PdfReader

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}


def read_resume(path: str) -> str:
    file = Path(path)
    if not file.exists():
        raise FileNotFoundError(f"Resume not found: {path}")

    ext = file.suffix.lower()
    if ext == ".pdf":
        reader = PdfReader(str(file))
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
    elif ext == ".docx":
        doc = Document(str(file))
        parts = [p.text for p in doc.paragraphs]
        # Many resumes keep skills/projects inside tables
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        text = "\n".join(parts)
    elif ext in (".txt", ".md"):
        text = file.read_text(encoding="utf-8", errors="ignore")
    else:
        raise ValueError(f"Unsupported file type '{ext}'. Use: {', '.join(sorted(SUPPORTED_EXTENSIONS))}")

    text = text.strip()
    if not text:
        raise ValueError(f"No text could be read from {path} (scanned image PDF?)")
    return text
